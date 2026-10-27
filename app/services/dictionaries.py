"""Словари: импорт, верификация разборов, правка."""

from __future__ import annotations

import csv
import io
import re
import unicodedata

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Dictionary, DictionaryEntry, TokenRow
from app.nlp.morph import parse_options
from app.nlp.normalize import decode_bytes
from app.nlp.tree import norm
from app.services.jobs import enqueue

# count_names — считать слова словаря и внутри имён собственных («любовь» в «Любовь Андреевна»); по умолчанию нет;
# no_categories — словарь без категорий: анализ показывает простой список слов; категории записей стираются (clear_categories)
DEFAULT_SETTINGS = {"include_participles": False, "include_gerunds": False, "merge_propn": True, "count_names": False,
                    "no_categories": False}
# невидимые символы и знаки ударения, которые попадают в словари из Word и сайтов: «благоуха́ния», «духи\u200b»
_INVISIBLE = dict.fromkeys(map(ord, "\u0301\u0300\u00ad\u200b\u200c\u200d\u2060\ufeff"), None)


def clean_surface(s: str) -> str:
    """Написание слова словаря без невидимых символов, ударений и лишних пробелов."""
    s = unicodedata.normalize("NFC", unicodedata.normalize("NFD", s or "").translate(_INVISIBLE))
    return re.sub(r"\s+", " ", s.replace("\u00a0", " ")).strip()


def entry_key(e: DictionaryEntry) -> tuple:
    """Записи с одним ключом находят одни и те же слова текста: «благоухание» и «благоухания» — лемма «благоухание»."""
    if e.phrase:
        return ("phrase", norm(clean_surface(e.surface)))
    return ("word", e.lemma_norm, e.upos or "", e.match_mode or "lemma")


def no_categories(d: Dictionary | None) -> bool:
    return bool(d is not None and (d.settings or {}).get("no_categories"))


def clear_categories(d: Dictionary) -> int:
    """Режим «без категорий»: стирает категории у всех слов словаря. Возвращает, у скольких слов они были."""
    n = 0
    for e in d.entries:
        if e.category:
            e.category = None
            n += 1
    return n

HEADER_WORDS = {"слово", "word", "лемма", "lemma", "токен", "token"}


def parse_dictionary_file(filename: str, data: bytes) -> list[dict]:
    """CSV/TSV/TXT/XLSX → [{surface, category, upos, comment}]."""
    name = filename.lower()
    rows: list[list[str]] = []
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True)
        for r in wb.active.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in r])
    else:
        content = decode_bytes(data)
        return parse_dictionary_text(content)
    return _rows_to_entries(rows)


def parse_dictionary_text(content: str) -> list[dict]:
    lines = [ln for ln in content.splitlines() if ln.strip() and not ln.lstrip().startswith("#")]
    if not lines:
        return []
    sample = "\n".join(lines[:20])
    if any(sep in sample for sep in (";", "\t")) or (lines[0].count(",") >= 1 and len(lines) > 1 and all(ln.count(",") <= 3 for ln in lines[:20])):
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";\t,")
            rows = list(csv.reader(lines, dialect))
            if rows and max(len(r) for r in rows) > 1:
                return _rows_to_entries(rows)
        except csv.Error:
            pass
    # простой список: по слову в строке или через запятую
    words: list[str] = []
    for ln in lines:
        words += [w for w in re.split(r"[,;]", ln)]
    return _rows_to_entries([[w] for w in words])


# колонки файла словаря по заголовку (иначе — по порядку: слово; категория; часть речи; комментарий)
COLUMN_NAMES = {"surface": {"слово", "word", "лемма", "lemma", "токен", "token"}, "category": {"категория", "category"},
                "upos": {"часть речи", "pos", "upos"}, "comment": {"комментарий", "comment"}}


def _rows_to_entries(rows: list[list[str]]) -> list[dict]:
    cols = {"surface": 0, "category": 1, "upos": 2, "comment": 3}
    if rows and rows[0] and rows[0][0].strip().lower() in HEADER_WORDS:
        head = [c.strip().lower() for c in rows[0]]
        cols = {k: next((i for i, h in enumerate(head) if h in names), None) for k, names in COLUMN_NAMES.items()}
        rows = rows[1:]
    out, seen = [], set()
    for r in rows:
        cells = [c.strip() for c in r]
        get = lambda k: cells[cols[k]] if cols[k] is not None and cols[k] < len(cells) else ""  # noqa: E731
        surface = get("surface")
        if not surface or surface.lower() in seen:
            continue
        seen.add(surface.lower())
        out.append({"surface": surface, "category": get("category") or None, "upos": get("upos").upper() or None,
                    "comment": get("comment") or None})
    return out


def make_entry(item: dict) -> DictionaryEntry:
    from app.services import phrases as P

    surface = clean_surface(item["surface"])
    phrase = None
    head_word = surface.split()[0] if surface else surface
    if P.is_phrase(surface):
        # словосочетание: ищется по главному слову, остальные слова проверяются при сопоставлении
        slots = P.parse(surface)
        head = P.head_slot(slots)
        phrase = {"slots": slots, "head": head}
        head_word = [w for w in P.split_words(surface) if w != P.GAP][sum(1 for x in slots[:head] if "gap" not in x)]
    opts = parse_options(head_word)
    wanted_upos = item.get("upos")
    best = None
    if wanted_upos:
        best = next((o for o in opts if o.upos == wanted_upos), None)
    best = best or (opts[0] if opts else None)
    lemma = best.lemma if best else head_word.lower()
    upos = best.upos if best else wanted_upos
    lemmas = {o.lemma for o in opts if o.score >= 0.05}
    pos_set = {o.upos for o in opts if o.score >= 0.05}
    return DictionaryEntry(
        surface=surface,
        lemma=lemma,
        lemma_norm=norm(lemma),
        upos=upos,
        match_mode="lemma",
        category=item.get("category"),
        comment=item.get("comment"),
        verified=bool(wanted_upos),
        ambiguous=len(lemmas) > 1 or len(pos_set) > 1,
        options=[{"lemma": o.lemma, "upos": o.upos, "tag": o.tag, "score": round(o.score, 3)} for o in opts[:6]],
        phrase=phrase,
    )


def create_dictionary(session: Session, name: str, items: list[dict], description: str | None = None,
                      settings: dict | None = None, enqueue_match: bool = True, skipped: list | None = None) -> Dictionary:
    """skipped — сюда попадают повторы: написание, совпавшее с уже добавленным словом по лемме («благоухания»)."""
    d = Dictionary(name=name.strip(), description=description, settings={**DEFAULT_SETTINGS, **(settings or {})})
    seen: dict[tuple, DictionaryEntry] = {}
    for it in items:
        e = make_entry(it)
        if _merge_into(seen, e, skipped):
            continue
        d.entries.append(e)
    session.add(d)
    session.flush()
    if enqueue_match:
        enqueue(session, "match", {"dictionary_id": d.id}, title=f"Сопоставление словаря «{d.name}»")
    return d


def _merge_into(seen: dict, e: DictionaryEntry, skipped: list | None) -> bool:
    """True — запись повторяет уже добавленную (та же лемма): её форма дописывается в комментарий, запись не создаётся."""
    k = entry_key(e)
    first = seen.get(k)
    if first is None:
        seen[k] = e
        return False
    if e.surface.lower() != first.surface.lower():
        _note_form(first, e.surface)
    if first.category is None and e.category:
        first.category = e.category
    if skipped is not None:
        skipped.append((e.surface, first.surface))
    return True


def _note_form(e: DictionaryEntry, form: str) -> None:
    note = f"также: {form}"
    if note not in (e.comment or ""):
        e.comment = f"{e.comment}; {note}" if e.comment else note


def add_entries(session: Session, d: Dictionary, items: list[dict], skipped: list | None = None) -> list[DictionaryEntry]:
    seen = {entry_key(e): e for e in d.entries}
    new = []
    for it in items:
        e = make_entry(it)
        if _merge_into(seen, e, skipped):
            continue
        d.entries.append(e)
        new.append(e)
    session.flush()
    if new:
        enqueue(session, "match", {"dictionary_id": d.id, "entry_ids": [e.id for e in new]},
                title=f"Сопоставление новых слов словаря «{d.name}»")
    return new


def update_entry(session: Session, entry: DictionaryEntry, *, lemma: str | None = None, upos: str | None = None,
                 match_mode: str | None = None, category: str | None = None) -> bool:
    """Возвращает True, если поменялся разбор или режим (тогда вхождения слова пересчитываются)."""
    changed = False
    if lemma is not None and lemma.strip() and lemma.strip() != entry.lemma:
        entry.lemma = lemma.strip()
        entry.lemma_norm = norm(entry.lemma)
        changed = True
    if upos is not None and (upos or None) != entry.upos:
        entry.upos = upos or None
        changed = True
    if match_mode in ("lemma", "forms") and match_mode != entry.match_mode:
        entry.match_mode = match_mode
        changed = True
    if category is not None:
        entry.category = category or None
    entry.verified = True
    session.flush()
    if changed:
        enqueue(session, "match", {"dictionary_id": entry.dictionary_id, "entry_ids": [entry.id]},
                title=f"Пересчёт слова «{entry.surface}»")
    return changed


def option_counts(session: Session, entry: DictionaryEntry) -> dict[str, int]:
    """Сколько токенов в базе у каждого варианта разбора (для экрана верификации)."""
    out = {}
    for o in entry.options or []:
        q = select(func.count()).select_from(TokenRow).where(TokenRow.lemma_norm == norm(o["lemma"]))
        if o.get("upos") and o["upos"] != "X":
            q = q.where(TokenRow.upos == o["upos"])
        out[f"{o['lemma']}|{o['upos']}"] = session.scalar(q) or 0
    return out


# ------------------------------------------------------------------ область словаря: все тексты или выбранные корпуса


def scope_corpora(d: Dictionary) -> list[int]:
    """Корпуса, с которыми сопоставляется словарь (settings["corpus_ids"]); пусто — все тексты."""
    return [int(x) for x in (d.settings or {}).get("corpus_ids", []) if str(x).isdigit()]


def scope_text_ids(session: Session, d: Dictionary) -> set[int] | None:
    """Тексты области словаря; None — все тексты."""
    from app.models import corpus_texts

    cids = scope_corpora(d)
    if not cids:
        return None
    return set(session.scalars(select(corpus_texts.c.text_id).where(corpus_texts.c.corpus_id.in_(cids))))


def matched_texts(session: Session, d: Dictionary) -> dict:
    """На скольких текстах посчитан словарь: {n — посчитано, total — текстов в области словаря, exact — точно ли}.

    Точное число ведётся с версии 0.14.1 (settings["matched_texts"], обновляет matching.run_match); у словарей,
    пересчитанных раньше, — число готовых текстов области (после полного пересчёта они все посчитаны)."""
    from app.models import TextDoc

    scope = scope_text_ids(session, d)
    q = select(TextDoc.id).where(TextDoc.status == "ready")
    if scope is not None:
        q = q.where(TextDoc.id.in_(scope or [-1]))
    ready = set(session.scalars(q))
    st = d.settings or {}
    if "matched_texts" in st:
        return {"n": len(ready & set(st["matched_texts"])), "total": len(ready), "exact": True}
    return {"n": len(ready) if d.matched_at else 0, "total": len(ready), "exact": False}


def sync_scope(session: Session, text_ids: list[int]) -> int:
    """После изменения корпусов у текстов: тексты, вошедшие в область словаря, ставятся на сопоставление,
    вышедшие — теряют вхождения этого словаря. Возвращает число поставленных задач."""
    from app.models import Hit, TextDoc

    if not text_ids:
        return 0
    n = 0
    for d in session.scalars(select(Dictionary)):
        scope = scope_text_ids(session, d)
        if scope is None:
            continue
        have = set(session.scalars(select(Hit.text_id).where(Hit.dictionary_id == d.id, Hit.text_id.in_(text_ids)).distinct()))
        ready = set(session.scalars(select(TextDoc.id).where(TextDoc.id.in_(text_ids), TextDoc.status == "ready")))
        add = sorted((ready & scope) - have)
        drop = sorted(have - scope)
        if drop:
            session.execute(Hit.__table__.delete().where(Hit.dictionary_id == d.id, Hit.text_id.in_(drop)))
            st = dict(d.settings or {})
            if "matched_texts" in st:
                d.settings = {**st, "matched_texts": sorted(set(st["matched_texts"]) - set(drop))}
        if add:
            enqueue(session, "match", {"dictionary_id": d.id, "text_ids": add},
                    title=f"Словарь «{d.name}» × новые тексты области ({len(add)})")
            n += 1
    return n


def duplicate_groups(d: Dictionary) -> list[list[DictionaryEntry]]:
    """Группы записей словаря с одной леммой (повторы и словоформы одного слова)."""
    groups: dict[tuple, list[DictionaryEntry]] = {}
    for e in sorted(d.entries, key=lambda e: e.id or 0):
        groups.setdefault(entry_key(e), []).append(e)
    return [g for g in groups.values() if len(g) > 1]


def merge_duplicates(session: Session, d: Dictionary) -> int:
    """Оставляет в каждой группе повторов одну запись (проверенную ✓, иначе первую): формы — в комментарий, категория —
    первая непустая. Вхождения удалённых записей совпадают с вхождениями оставшейся и удаляются вместе с ними."""
    removed = 0
    for g in duplicate_groups(d):
        keep = next((e for e in g if e.verified), g[0])
        for e in g:
            if e is keep:
                continue
            if e.surface.lower() != keep.surface.lower():
                _note_form(keep, e.surface)
            if not keep.category and e.category:
                keep.category = e.category
            d.entries.remove(e)
            session.delete(e)
            removed += 1
    if removed:
        for e in d.entries:  # написание без невидимых символов
            e.surface = clean_surface(e.surface)
    session.flush()
    return removed
