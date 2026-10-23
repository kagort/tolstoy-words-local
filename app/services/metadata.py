"""Таблица метаданных текстов (CSV / XLSX): какой файл — какое название, автор, год, жанр."""

from __future__ import annotations

import csv
import io
import re
from pathlib import PurePath

from app.nlp.normalize import decode_bytes

ALIASES = {
    "file": ("файл", "file", "filename", "имя файла", "textid", "id"),
    "title": ("название", "заглавие", "title", "texttitle"),
    "short_title": ("короткое", "короткое название", "краткое название", "сокращение", "short", "short_title", "shorttitle"),
    "author": ("автор", "author", "text_author"),
    "years": ("год", "годы", "year", "years", "text_year_creation", "дата"),
    "genre": ("жанр", "genre", "text_genre"),
    "source_type": ("тип", "type", "source_type", "тип текста"),
    "generator": ("модель", "генератор", "generator", "промпт", "prompt"),
    "corpus": ("корпус", "корпуса", "подкорпус", "corpus", "corpora", "collection"),
}
# значения «нет данных» в таблицах: прочерки, пустые маркеры
EMPTY_VALUES = {"-", "—", "–", "--", "нет", "n/a", "na", "nan", "none", "null", "?"}

SOURCE_TYPE_WORDS = {
    "natural": ("natural", "литература", "естественный", "человек", "оригинал"),
    "generated": ("generated", "нейросеть", "сгенерирован", "сгенерированный", "llm", "ии", "ai", "gpt", "модель"),
}


class _Semicolon(csv.excel):
    delimiter = ";"


def normalize_source_type(value: str | None) -> str | None:
    if not value:
        return None
    v = value.strip().lower()
    for key, words in SOURCE_TYPE_WORDS.items():
        if any(v.startswith(w) for w in words):
            return key
    return "other"


def parse_years(value: str | None) -> tuple[int | None, int | None]:
    ys = [int(y) for y in re.findall(r"\d{3,4}", value or "")]
    return (ys[0], ys[-1]) if ys else (None, None)


def _canonical(header: str) -> str | None:
    h = header.strip().lower()
    for key, names in ALIASES.items():
        if h in names:
            return key
    return None


def _num(key: str) -> str | None:
    """«007» → «7»: числовые имена файлов («1.txt», «001.txt») совпадают с ID в таблице."""
    return (key.lstrip("0") or "0") if key.isdigit() else None


def file_keys(filename: str) -> list[str]:
    """По каким ключам искать строку метаданных для файла: имя с расширением, без расширения, число без нулей."""
    name = PurePath(filename).name.lower()
    stem = PurePath(name).stem
    keys = [name, stem]
    if _num(stem):
        keys.append(_num(stem))
    return keys


def _clean(value: str) -> str:
    v = (value or "").strip()
    return "" if v.lower() in EMPTY_VALUES else v


def parse_meta_table(filename: str, data: bytes) -> dict[str, dict]:
    """→ {ключ файла (нижний регистр): {title, author, years, genre, source_type, generator}}"""
    rows: list[list[str]] = []
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True)
        rows = [["" if v is None else str(v) for v in r] for r in wb.active.iter_rows(values_only=True)]
    else:
        text = decode_bytes(data)
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if not lines:
            return {}
        try:
            dialect = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=";,\t")
        except csv.Error:
            dialect = _Semicolon
        rows = list(csv.reader(lines, dialect))
    if not rows:
        return {}
    header = [_canonical(h) for h in rows[0]]
    if "file" not in header:
        raise ValueError("В таблице метаданных нет колонки «файл» (имя файла, к которому относится строка)")
    out: dict[str, dict] = {}
    for r in rows[1:]:
        rec = {k: _clean(r[i]) if i < len(r) else "" for i, k in enumerate(header) if k}
        key = rec.pop("file", "").lower()
        if not key:
            continue
        rec["_key"] = key  # исходный ID из таблицы — сохраняется в сведениях текста
        out[PurePath(key).name] = rec
        out.setdefault(PurePath(key).stem, rec)
        if _num(PurePath(key).stem):
            out.setdefault(_num(PurePath(key).stem), rec)
    return out


def lookup(meta: dict[str, dict], filename: str) -> dict:
    for k in file_keys(filename):
        if k in meta:
            return meta[k]
    return {}


def unmatched_rows(meta: dict[str, dict], used: set[int]) -> list[str]:
    """ID строк таблицы, для которых не нашлось файла (used — id() использованных записей)."""
    seen, out = set(), []
    for rec in meta.values():
        if id(rec) in seen:
            continue
        seen.add(id(rec))
        if id(rec) not in used:
            out.append(rec.get("_key", "?"))
    return out


def split_corpora(value: str) -> list[str]:
    """Колонка «корпус»: несколько корпусов через «;» или «|» (запятая — допустима внутри названия)."""
    parts = re.split(r"[;|]", value or "")
    return [p.strip() for p in parts if p.strip() and p.strip().lower() not in EMPTY_VALUES]
